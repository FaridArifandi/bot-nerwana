import os
import sqlite3
import openpyxl

from database import get_connection, log_audit

EXCEL_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "review_usulan.xlsx")

def sync_excel_to_db():
    if not os.path.exists(EXCEL_PATH):
        print(f"[ERROR] File {EXCEL_PATH} tidak ditemukan.")
        return 0

    print(f"Membaca file {EXCEL_PATH}...")
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb["Review Usulan Flag 3"]

    conn = get_connection()
    cursor = conn.cursor()

    approved_count = 0
    updated_count = 0

    header_row = [str(cell or '').strip().upper() for cell in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    col_map = {name: idx for idx, name in enumerate(header_row)}
    
    col_aid = col_map.get("ASSIGNMENT_ID", 1)
    col_nama = col_map.get("NAMA_USAHA", 2)
    col_pbj = col_map.get("PBJ_USULAN", 10)
    col_prod = col_map.get("PROD_USULAN", 11)
    col_ops = col_map.get("OPS_USULAN", 12)
    col_rasio = col_map.get("RASIO_BARU", 14)
    col_cat = col_map.get("KODE_CATATAN", 16)
    col_alasan = col_map.get("ALASAN_USULAN", 17)
    col_app = col_map.get("APPROVED", 19)
    col_rev = col_map.get("CATATAN_REVIEWER", 20)

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        aid = str(row[col_aid] or '').strip() if col_aid < len(row) else ''
        nama = str(row[col_nama] or '').strip() if col_nama < len(row) else ''
        if not aid or not nama:
            continue

        # Data usulan dari Excel (jika user ada mengedit di Excel)
        pbj_usulan = float(row[col_pbj] or 0) if col_pbj < len(row) else 0
        prod_usulan = float(row[col_prod] or 0) if col_prod < len(row) else 0
        ops_usulan = float(row[col_ops] or 0) if col_ops < len(row) else 0
        rasio_baru = float(row[col_rasio] or 0) if col_rasio < len(row) else 0
        kode_catatan = str(row[col_cat] or '').strip() if col_cat < len(row) else ''
        alasan = str(row[col_alasan] or '').strip() if col_alasan < len(row) else ''
        approved_val = str(row[col_app] or '').strip().upper() if col_app < len(row) else ''
        catatan_reviewer = str(row[col_rev] or '').strip() if col_rev < len(row) else ''

        # Tentukan status baru
        if approved_val == 'Y':
            new_status = 'APPROVED'
            approved_count += 1
        else:
            new_status = 'PENDING'

        # Cek apakah status di DB sudah APPLIED / SUBMITTED / COMPLETED / MARKED (jangan ditimpa ke APPROVED jika sudah selesai)
        cur = cursor.execute("SELECT status, tindak_lanjut FROM kasus_nerwana WHERE assignment_id = ? AND nama_usaha = ?", (aid, nama)).fetchone()
        if cur and (cur['status'] in ['APPLIED', 'SUBMITTED', 'MARKED', 'COMPLETED'] or cur['tindak_lanjut'] == 'sudah_direvisi'):
            continue

        cursor.execute("""
        UPDATE kasus_nerwana SET
            nilai_baru_pbj = ?,
            nilai_baru_prod = ?,
            nilai_baru_ops = ?,
            rasio_baru = ?,
            kode_catatan = ?,
            alasan_usulan = ?,
            status = ?,
            alasan_review = ?,
            updated_at = CURRENT_TIMESTAMP
        WHERE assignment_id = ? AND nama_usaha = ?
        """, (
            pbj_usulan, prod_usulan, ops_usulan, rasio_baru, kode_catatan, alasan,
            new_status, catatan_reviewer, aid, nama
        ))
        updated_count += 1

    conn.commit()
    conn.close()
    wb.close()

    print(f"[SYNC] Selesai menyinkronkan {updated_count} baris dari Excel.")
    print(f"[SYNC] Total kasus siap dieksekusi (status APPROVED): {approved_count} baris.")
    return approved_count

if __name__ == "__main__":
    sync_excel_to_db()
